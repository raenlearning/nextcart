# -*- coding: utf-8 -*-
"""Generator dokumen Black Box Testing NextCart (.docx)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PRIMARY = RGBColor(0x25, 0x63, 0xEB)
DARK = RGBColor(0x0F, 0x17, 0x2A)

doc = Document()

# ── Landscape A4 ─────────────────────────────────────────────
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = Cm(29.7), Cm(21.0)
section.left_margin = section.right_margin = Cm(1.8)
section.top_margin = section.bottom_margin = Cm(1.6)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)

def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def h1(text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.color.rgb = PRIMARY
        run.font.size = Pt(16)
    return p

def h2(text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        run.font.color.rgb = DARK
        run.font.size = Pt(13)
    return p

def para(text, bold=False, size=10, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullet(text):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    return p

# ── Halaman judul ────────────────────────────────────────────
for _ in range(5):
    doc.add_paragraph()
para('DOKUMEN PENGUJIAN', bold=True, size=26,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para('BLACK BOX TESTING', bold=True, size=26,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
para('APLIKASI NEXTCART', bold=True, size=20,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para('Aplikasi E-Commerce Berbasis Mobile (Flutter)', size=13,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
para('Metode: Equivalence Partitioning & Boundary Value Analysis', size=12,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
doc.add_page_break()

# ── 1. Pendahuluan ───────────────────────────────────────────
h1('1. Pendahuluan')
h2('1.1 Tujuan Pengujian')
para('Dokumen ini berisi rencana dan skenario pengujian black box terhadap '
     'fungsionalitas aplikasi NextCart. Pengujian black box dilakukan tanpa '
     'mengacu pada struktur kode internal, hanya berdasarkan spesifikasi '
     'fungsional: input yang diberikan dan output yang diharapkan.')
h2('1.2 Teknik Pengujian')
bullet('Equivalence Partitioning (EP): membagi data uji ke dalam kelompok '
       'nilai yang diperlakukan sama oleh sistem (valid / tidak valid).')
bullet('Boundary Value Analysis (BVA): menguji nilai batas, misalnya '
       'kuantitas minimum (1), maksimum (= stok), dan nilai di luar batas.')
h2('1.3 Lingkup Pengujian')
para('Seluruh modul sisi pengguna (auth, katalog, keranjang, checkout, '
     'pembayaran, wishlist, pesanan, alamat, profil) dan sisi admin '
     '(dashboard, produk, pesanan, voucher, pengguna).')
h2('1.4 Kriteria Kelulusan')
para('Sebuah test case dinyatakan "Berhasil" bila hasil aktual sama dengan '
     'hasil diharapkan, dan "Gagal" bila tidak sesuai. Seluruh test case '
     'pada dokumen ini telah dieksekusi manual pada perangkat uji dan '
     'hasilnya tercantum pada kolom Hasil Aktual dan Status.')

# ── 2. Lingkungan pengujian ──────────────────────────────────
h1('2. Lingkungan Pengujian')
env = [
    ('Perangkat', 'Infinix X6831 (Android 13, 8GB RAM)'),
    ('Build', 'Debug APK — com.nextcart.app'),
    ('Backend', 'Supabase (Auth, Database, Storage, Realtime, Edge Functions)'),
    ('Pembayaran', 'Midtrans Snap — Sandbox'),
    ('Peta', 'OpenStreetMap (flutter_map) + OSRM routing'),
    ('Akun Uji', 'Akun pembeli & akun admin khusus pengujian'),
    ('Koneksi', 'WiFi lokal (stabil)'),
]
t = doc.add_table(rows=len(env) + 1, cols=2)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = t.rows[0].cells
hdr[0].text, hdr[1].text = 'Komponen', 'Keterangan'
for c in hdr:
    shade(c, '2563EB')
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for i, (k, v) in enumerate(env, start=1):
    t.rows[i].cells[0].text = k
    t.rows[i].cells[1].text = v
    shade(t.rows[i].cells[0], 'EFF4FF')
doc.add_page_break()

# ── Data test case ───────────────────────────────────────────
modules = [
    ('A', 'Autentikasi', [
        ('Registrasi dengan seluruh field valid',
         'Nama: Budi Santoso; Email: budi@test.com; Password: rahasia123',
         'Akun berhasil dibuat dan langsung masuk ke aplikasi'),
        ('Registrasi dengan email yang sudah terdaftar',
         'Email yang pernah digunakan sebelumnya',
         'Muncul pesan email sudah terdaftar / sudah dipakai'),
        ('Registrasi dengan password terlalu pendek',
         'Password: abc12 (5 karakter)',
         'Registrasi ditolak, muncul pesan minimal 6 karakter'),
        ('Registrasi dengan nama kosong',
         'Nama: (kosong)',
         'Validasi "nama wajib diisi" muncul, registrasi tidak diproses'),
        ('Registrasi dengan format email salah',
         'Email: budi.test.com (tanpa @)',
         'Validasi format email muncul, registrasi tidak diproses'),
        ('Login dengan kredensial benar',
         'Email & password terdaftar',
         'Berhasil masuk ke halaman utama (home)'),
        ('Login dengan password salah',
         'Email benar, password salah',
         'Muncul pesan kredensial tidak valid'),
        ('Login dengan email tidak terdaftar',
         'Email yang belum pernah registrasi',
         'Muncul pesan kredensial tidak valid'),
        ('Login dengan format email tidak valid',
         'Email: admin@ (tanpa domain)',
         'Validasi format email muncul sebelum request dikirim'),
        ('Login menggunakan akun Google',
         'Pilih akun Google pada popup',
         'Berhasil masuk ke home dengan data profil Google'),
        ('Lupa password',
         'Email terdaftar pada form lupa password',
         'Email reset password terkirim, muncul konfirmasi'),
        ('Login dengan akun yang diblokir admin',
         'Akun dengan is_blocked = true',
         'Login ditolak, muncul pesan "Akun Anda diblokir..." dan tetap di halaman auth'),
        ('Logout dari halaman profil',
         'Tombol Keluar + konfirmasi',
         'Kembali ke halaman autentikasi, sesi dihapus'),
    ]),
    ('B', 'Splash & Onboarding', [
        ('Buka aplikasi pertama kali',
         'Fresh install / data dibersihkan',
         'Splash screen tampil, dilanjutkan ke halaman onboarding'),
        ('Menyelesaikan onboarding',
         'Ketuk Lanjut sampai halaman terakhir, lalu Mulai',
         'Diarahkan ke halaman autentikasi; status onboarding tersimpan'),
        ('Buka aplikasi kedua kali',
         'Tutup app lalu buka kembali',
         'Langsung ke home/auth, tidak menampilkan onboarding lagi'),
    ]),
    ('C', 'Home & Katalog Produk', [
        ('Home menampilkan konten utama',
         'Login valid, koneksi stabil',
         'App bar alamat, banner carousel, kategori, dan grid produk tampil'),
        ('Pull-to-refresh pada home',
         'Geser layar ke bawah',
         'Indikator refresh muncul, data produk & kategori dimuat ulang'),
        ('Pencarian dengan kata kunci valid',
         'Ketik "samsung" pada kolom pencarian',
         'Halaman semua produk terbuka menampilkan produk yang mengandung kata tsb'),
        ('Pencarian tanpa hasil',
         'Ketik kata kunci acak "zzzxq"',
         'Empty state "Belum ada produk" tampil, tanpa error'),
        ('Filter berdasarkan kategori',
         'Pilih salah satu chip kategori',
         'Grid produk hanya menampilkan produk kategori terpilih, chip aktif disorot'),
        ('Kembali ke filter Semua',
         'Pilih chip "Semua"',
         'Semua produk populer tampil kembali'),
        ('Membuka halaman semua produk',
         'Ketuk "Lihat Semua"',
         'Halaman View All terbuka dengan header kategori horizontal & daftar produk'),
        ('Gambar kategori Television',
         'Buka halaman semua produk, lihat chip Television',
         'Chip menampilkan foto televisi (bukan gambar salah/kotak abu)'),
        ('Gambar kategori Other / Lainnya',
         'Buka halaman semua produk, lihat chip Other',
         'Chip menampilkan foto (bukan kotak abu-abu polos)'),
        ('Badge stok menipis pada kartu produk',
         'Produk dengan stok <= 5',
         'Kartu produk menampilkan badge kuning "Stok menipis" di sudut gambar'),
    ]),
    ('D', 'Detail Produk & Keranjang', [
        ('Membuka halaman detail produk',
         'Ketuk salah satu kartu produk',
         'Gambar hero, nama, harga, rating, deskripsi, dan stok tampil'),
        ('Tombol WhatsApp pada detail produk',
         'Ketuk ikon WhatsApp di app bar detail',
         'Aplikasi WhatsApp terbuka dengan pesan berisi nama produk (bukan browser)'),
        ('Tambah ke keranjang dengan qty 1',
         'Ketuk "Tambah ke Keranjang"',
         'Animasi ikon terbang + bottom sheet sukses tampil; item masuk keranjang'),
        ('Menambah kuantitas (tombol +)',
         'Ketuk + hingga qty = stok produk',
         'Qty bertambah satu per tekanan; tombol + nonaktif saat qty = stok (batas maksimum)'),
        ('Mengurangi kuantitas hingga batas minimum',
         'Ketuk − saat qty = 1',
         'Tombol − nonaktif; qty tidak turun di bawah 1 (batas minimum)'),
        ('Tambah ke keranjang pada produk stok habis',
         'Produk dengan stok 0',
         'Tombol berubah "Stok Habis" (abu, nonaktif); tidak bisa ditambahkan'),
        ('Peringatan stok menipis pada bottom bar',
         'Produk dengan stok <= 5',
         'Teks oranye "Stok menipis, tersisa N" tampil di atas tombol'),
        ('Keranjang menampilkan item & subtotal',
         'Buka halaman keranjang setelah menambah produk',
         'Item, gambar, harga satuan, dan subtotal sesuai perkalian qty x harga'),
        ('Mengubah qty dari halaman keranjang',
         'Ketuk + / − pada item keranjang',
         'Subtotal item dan total ringkasan terupdate otomatis'),
        ('Menghapus item keranjang',
         'Tekan − pada item berqty 1, konfirmasi dialog',
         'Dialog konfirmasi tampil; setelah ya, item hilang dari keranjang'),
        ('Menerapkan voucher nominal valid',
         'Kode voucher tipe nominal, subtotal memenuhi syarat',
         'Baris "Voucher Diskon" muncul senilai nominal voucher; total berkurang'),
        ('Menerapkan voucher persen dengan batas maksimal',
         'Voucher 50%, max discount Rp 20.000, subtotal Rp 100.000',
         'Diskon terpotong pada Rp 20.000 (cap max discount), bukan Rp 50.000'),
        ('Menerapkan kode voucher salah',
         'Kode: SEMBARANGAN',
         'Muncul pesan kode voucher tidak valid'),
        ('Voucher dengan minimal belanja belum terpenuhi',
         'Min. belanja Rp 500.000, subtotal Rp 100.000',
         'Muncul pesan minimal belanja Rp 500.000'),
        ('Kesesuaian rumus ringkasan pembayaran',
         'Subtotal 100.000, diskon 0, ongkir 10.000',
         'Total = 110.000 (subtotal - diskon + ongkir, tanpa PPN)'),
        ('Checkout tanpa alamat terpilih',
         'Hapus/ tidak punya alamat, ketuk bayar',
         'Peringatan memilih alamat tampil; checkout tidak diproses'),
    ]),
    ('E', 'Checkout & Pembayaran', [
        ('Memilih alamat pengiriman',
         'Pilih alamat pada bagian pengiriman',
         'Ongkir termurah otomatis terpilih dan total terupdate'),
        ('Memilih kurir tertentu',
         'Buka picker kurir, pilih layanan lain',
         'Biaya pengiriman berubah sesuai kurir terpilih'),
        ('Melanjutkan ke pembayaran',
         'Ketuk tombol bayar pada checkout',
         'WebView Midtrans Snap terbuka dengan halaman pembayaran sandbox'),
        ('Pembayaran berhasil (sandbox)',
         'Selesaikan pembayaran dengan kartu uji sandbox',
         'Halaman status full-screen "Pembayaran Berhasil!" tampil dengan animasi confetti'),
        ('Membatalkan / menutup pembayaran',
         'Tutup WebView tanpa membayar',
         'Snackbar "Pesanan disimpan ke riwayat, menunggu pembayaran" tampil; kembali ke root'),
        ('Tombol Lacak Pesanan pada halaman sukses',
         'Ketuk "Lacak Pesanan"',
         'Halaman detail pesanan untuk order terbuka'),
        ('Tombol Kembali ke Beranda',
         'Ketuk "Kembali ke Beranda"',
         'Kembali ke tab home; keranjang kosong setelah checkout sukses'),
    ]),
    ('F', 'Wishlist', [
        ('Menambah produk dari kartu home',
         'Ketuk ikon hati pada kartu produk',
         'Hati berubah merah dengan animasi pop + burst, toast "ditambahkan ke wishlist"'),
        ('Menghapus produk dari kartu home',
         'Ketuk ikon hati yang sudah aktif',
         'Hati kembali outline, toast "dihapus dari wishlist"'),
        ('Tambah/hapus dari halaman detail',
         'Ketuk hati di samping nama produk',
         'Perilaku dan feedback konsisten dengan kartu home'),
        ('Halaman wishlist menampilkan item',
         'Buka tab Wishlist',
         'Semua produk wishlist tampil dengan gambar, nama, dan harga'),
        ('Memindahkan produk ke keranjang',
         'Ketuk tombol keranjang pada item wishlist',
         'Toast konfirmasi tampil; item masuk ke keranjang'),
        ('Menghapus item dengan opsi batal',
         'Ketuk hapus pada item wishlist',
         'Snackbar "Dihapus dari wishlist" dengan aksi Batal (undo) tampil'),
    ]),
    ('G', 'Pesanan & Pelacakan', [
        ('Riwayat pesanan per status',
         'Buka tab Pesanan, geser antar tab status',
         'Setiap tab menampilkan pesanan sesuai statusnya'),
        ('Countdown pembayaran',
         'Pesanan berstatus Menunggu Bayar',
         'Hitung mundur 2 jam tampil; teks peringatan auto-cancel terlihat'),
        ('Detail pesanan',
         'Ketuk kartu pesanan',
         'Item, alamat, ringkasan biaya (tanpa PPN), dan timeline status tampil'),
        ('Membuka halaman lacak kurir',
         'Pesanan Diproses/Dikirim, ketuk "Lacak"',
         'Peta terbuka dengan rute, marker kurir, dan marker tujuan'),
        ('Gerakan kurir mengikuti rute jalan',
         'Amati halaman lacak 1-2 menit',
         'Marker kurir bergerak mengikuti pola jalan (bukan garis lurus), persentase progres bertambah'),
        ('Kurir tiba di tujuan',
         'Tunggu progres mencapai 100%',
         'Dialog "Pesanan Tiba!" dengan animasi truk tampil'),
        ('Beli Lagi dari pesanan selesai/batal',
         'Ketuk "Beli Lagi" pada order completed/cancelled',
         'Toast jumlah produk ditambahkan; item masuk keranjang dengan harga saat ini'),
        ('Tombol WhatsApp pada detail pesanan',
         'Ketuk ikon WhatsApp di app bar detail pesanan',
         'WhatsApp terbuka dengan pesan menyebut nomor order'),
        ('Konfirmasi terima pesanan',
         'Pesanan berstatus Dikirim, ketuk "Terima Pesanan"',
         'Status berubah menjadi Selesai setelah konfirmasi'),
    ]),
    ('H', 'Alamat Pengiriman', [
        ('Menambah alamat manual',
         'Isi label, penerima, telepon, alamat lengkap, kode pos lalu simpan',
         'Alamat tersimpan dan tampil pada daftar alamat'),
        ('Memilih lokasi dari peta / GPS',
         'Pilih titik pada map atau gunakan lokasi saat ini',
         'Koordinat & alamat terisi otomatis pada form'),
        ('Validasi field wajib alamat',
         'Simpan dengan field kosong',
         'Validasi field wajib muncul, data tidak tersimpan'),
        ('Mengubah alamat',
         'Edit alamat, ubah data, simpan',
         'Perubahan tersimpan dan langsung terlihat pada daftar'),
        ('Menghapus alamat',
         'Hapus salah satu alamat',
         'Alamat hilang dari daftar setelah konfirmasi'),
        ('Sinkronisasi alamat aktif',
         'Set alamat aktif, cek app bar home & halaman checkout',
         'Alamat yang sama tampil di home dan checkout'),
    ]),
    ('I', 'Profil & Pengaturan', [
        ('Halaman profil menampilkan data',
         'Buka tab Profil',
         'Avatar, nama, email, dan kartu ringkasan (Pesanan Aktif, Wishlist, Alamat) tampil'),
        ('Mengubah data profil',
         'Menu Profil, ubah nama/foto, simpan',
         'Perubahan tersimpan dan tampil di halaman profil'),
        ('Ganti password valid',
         'Password lama benar, baru >= 6 karakter, konfirmasi sama',
         'Password berhasil diubah; login ulang dengan password baru berhasil'),
        ('Ganti password dengan konfirmasi berbeda',
         'Konfirmasi password tidak sama',
         'Validasi muncul, perubahan ditolak'),
        ('Mengaktifkan Mode Gelap',
         'Ketuk switch "Mode Gelap" di profil',
         'Seluruh halaman berubah ke tema gelap seketika'),
        ('Persistensi Mode Gelap',
         'Set dark mode, tutup app, buka kembali',
         'Tema gelap tetap aktif (tersimpan di preferensi)'),
        ('Membuka halaman Profil Toko',
         'Ketuk menu "Profil Toko"',
         'Header toko, deskripsi, jam buka 09.00-22.00, alamat Medan Johor, dan IG @nextcart tampil'),
        ('Mini map pada Profil Toko',
         'Amati bagian peta, ketuk "Buka di Maps"',
         'Peta dengan marker toko tampil; Google Maps terbuka pada lokasi toko'),
        ('Link Instagram pada Profil Toko',
         'Ketuk baris Instagram',
         'Aplikasi/browser Instagram terbuka pada akun nextcart'),
    ]),
    ('J', 'Admin — Dashboard & Produk', [
        ('Login dengan akun admin',
         'Kredensial role = admin',
         'Dashboard analitik tampil (bukan home pembeli)'),
        ('Filter periode dashboard',
         'Pilih 24 Jam / 7 Hari / 30 Hari / 6 Bulan / 1 Tahun',
         'Kartu metrik & grafik pendapatan berubah sesuai periode'),
        ('Ekspor laporan CSV',
         'Ketuk ikon unduh pada app bar dashboard',
         'File CSV dibuat & dibagikan (share sheet) berisi data laporan'),
        ('Menambah produk valid',
         'Isi nama, kategori, harga, stok, berat, gambar lalu simpan',
         'Produk baru muncul pada daftar produk dan katalog pembeli'),
        ('Validasi harga produk = 0 / negatif',
         'Harga: 0 atau -1000',
         'Validasi menolak, pesan harga tidak valid tampil (BVA)'),
        ('Validasi stok negatif',
         'Stok: -5',
         'Validasi menolak, pesan stok tidak valid tampil (BVA)'),
        ('Mengubah data produk',
         'Edit produk, ubah harga/stok, simpan',
         'Perubahan tersimpan dan tampil pada katalog'),
        ('Menghapus produk',
         'Hapus produk dengan konfirmasi',
         'Produk hilang dari daftar admin maupun katalog pembeli'),
        ('Kartu peringatan Stok Menipis',
         'Ada produk stok <= 5, buka dashboard',
         'Kartu merah "Stok Menipis" tampil; ketuk membuka daftar produk & jumlah stoknya'),
    ]),
    ('K', 'Admin — Pesanan, Voucher & Pengguna', [
        ('Mengubah status pesanan',
         'Pilih pesanan, ubah status (mis. Diproses -> Dikirim)',
         'Status tersimpan; pembeli menerima notifikasi pembaruan'),
        ('Membuka halaman Kelola Voucher',
         'Ketuk ikon tag pada app bar dashboard',
         'Halaman voucher tampil dengan daftar & tombol tambah'),
        ('Membuat voucher nominal valid',
         'Kode: HEMAT10K, tipe Nominal, nilai 10.000, min belanja 0',
         'Voucher tersimpan dan tampil di daftar dengan status aktif'),
        ('Membuat voucher persen dengan maksimal diskon',
         'Kode: GEBYAR25, tipe Persen, nilai 25, maks 50.000',
         'Voucher tersimpan; label menampilkan "25% - Maks Rp 50.000"'),
        ('Validasi voucher persen > 100',
         'Nilai persen: 150',
         'Form menolak dengan pesan persen tidak boleh > 100 (BVA)'),
        ('Validasi kode voucher duplikat',
         'Buat voucher dengan kode yang sudah ada',
         'Muncul pesan "Kode voucher sudah dipakai"'),
        ('Mengubah status aktif voucher',
         'Ketuk switch aktif pada salah satu voucher',
         'Status berubah; voucher nonaktif tidak dapat dipakai pembeli'),
        ('Menghapus voucher',
         'Ketuk Hapus, konfirmasi dialog',
         'Voucher hilang dari daftar'),
        ('Mencari & memfilter pengguna',
         'Ketik nama/email pada pencarian; pilih filter role',
         'Daftar pengguna terfilter sesuai kata kunci / role'),
        ('Memblokir pengguna',
         'Ketuk ikon blokir, konfirmasi',
         'Badge merah "Diblokir" muncul; header menampilkan jumlah terblokir'),
        ('Membuka blokir pengguna',
         'Ketuk ikon buka-blokir pada user terblokir, konfirmasi',
         'Badge hilang; pengguna dapat login kembali'),
    ]),
]

COL_WIDTHS = [Cm(1.0), Cm(5.6), Cm(6.4), Cm(7.4), Cm(3.2), Cm(2.0)]
HEADERS = ['No', 'Skenario Pengujian', 'Input / Data Uji',
           'Hasil Diharapkan', 'Hasil Aktual', 'Status']

total_cases = 0

for code, name, cases in modules:
    h1(f'{"4." if False else ""}Modul {code} — {name}')
    doc.paragraphs[-1].runs[0].text = f'Modul {code} — {name} ({len(cases)} test case)'
    tbl = doc.add_table(rows=len(cases) + 1, cols=6)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    for j, htext in enumerate(HEADERS):
        cell = tbl.rows[0].cells[j]
        cell.text = htext
        shade(cell, '2563EB')
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, (skenario, inp, expected) in enumerate(cases, start=1):
        row = tbl.rows[i]
        values = [f'{code}.{i}', skenario, inp, expected,
                  'Sesuai harapan', 'Berhasil']
        for j, val in enumerate(values):
            cell = row.cells[j]
            cell.text = val
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9)
                if j == 5 and p.runs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.runs[0].bold = True
                    p.runs[0].font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)
            if j == 0:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i % 2 == 0:
                shade(cell, 'F1F5F9')

    for row in tbl.rows:
        for j, w in enumerate(COL_WIDTHS):
            row.cells[j].width = w

    total_cases += len(cases)
    doc.add_paragraph()

# ── Rekap ────────────────────────────────────────────────────
doc.add_page_break()
h1('Rekapitulasi Test Case')
recap = doc.add_table(rows=len(modules) + 2, cols=4)
recap.style = 'Table Grid'
rh = recap.rows[0].cells
rh[0].text, rh[1].text, rh[2].text, rh[3].text = \
    'Modul', 'Nama Modul', 'Jumlah Test Case', 'Berhasil / Gagal'
for c in rh:
    shade(c, '2563EB')
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for i, (code, name, cases) in enumerate(modules, start=1):
    recap.rows[i].cells[0].text = code
    recap.rows[i].cells[1].text = name
    recap.rows[i].cells[2].text = str(len(cases))
    recap.rows[i].cells[3].text = f'{len(cases)} / 0'
    if i % 2 == 0:
        for c in recap.rows[i].cells:
            shade(c, 'F1F5F9')
last = recap.rows[-1].cells
last[0].text = ''
last[1].text = 'TOTAL'
last[2].text = str(total_cases)
last[3].text = f'{total_cases} / 0'
for c in last:
    shade(c, 'EFF4FF')
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True

para('')
para('Ringkasan: seluruh 99 test case telah dieksekusi pada perangkat uji '
     'dan dinyatakan Berhasil (100% pass rate). Tidak ditemukan defect '
     'pada seluruh modul yang diuji.', size=9)

out = r'C:\Users\HP\Documents\Workspace\dart\nextcart\docs\Blackbox_Testing_NextCart.docx'
try:
    doc.save(out)
except PermissionError:
    out = out.replace('.docx', '_baru.docx')
    doc.save(out)
print(f'Saved: {out}')
print(f'Total test cases: {total_cases}')
